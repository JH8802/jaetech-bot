"""엑셀/CSV 일괄 등록: 양식 만들기, 파일 읽기·검증, 현재 마스터 내보내기.

지원 형식: .xlsx / .xls / .csv
 - 시트 이름은 상관없음. 헤더 행에 '원료명'이 있으면 원료, '부자재명'이 있으면 부자재로 인식.
 - 헤더 위에 제목 줄이 있어도 됨 (앞 15행 안에서 헤더를 찾음).
 - 빈 칸은 '기존 값 유지' (신규면 기본값).
 - 오류가 하나라도 있으면 아무것도 등록하지 않는다.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .engine import CATEGORIES

CAT_BY_LABEL = {v: k for k, v in CATEGORIES.items()}
CAT_ALIASES = {
    "롤포장지": "roll", "롤": "roll", "롤필름": "roll", "roll": "roll",
    "포장지": "pack", "포장": "pack", "소포장지": "pack", "pack": "pack",
    "인케이스": "incase", "케이스": "incase", "ic": "incase", "incase": "incase",
    "카톤박스": "box", "카톤": "box", "박스": "box", "박스비": "box", "ct": "box", "box": "box",
    "소분비": "split", "소분": "split", "split": "split",
    "기타": "etc", "etc": "etc",
}

# 통화 별칭 → 통화 코드 (그 외에는 입력한 코드를 대문자로 사용)
CUR_ALIASES = {
    "usd": "USD", "미국": "USD", "달러": "USD", "미국달러": "USD", "$": "USD",
    "vnd": "VND", "베트남": "VND", "동": "VND", "베트남동": "VND",
    "try": "TRY", "튀르키예": "TRY", "터키": "TRY", "리라": "TRY", "튀르키예리라": "TRY", "터키리라": "TRY",
    "krw": "KRW", "원": "KRW", "원화": "KRW", "한국": "KRW", "대한민국": "KRW",
}


class Pct(float):
    """엑셀에서 '%' 서식 셀이었던 값. 이미 % 단위(5.0 = 5%)로 변환되어 있다."""


def _norm(s) -> str:
    return re.sub(r"[\s()%/·_\-\[\]]", "", str(s)).lower() if s is not None else ""


# 헤더 동의어 (정규화된 형태)
ING_NAME = {"원료명", "원료", "원재료", "원재료명"}
MAT_NAME = {"부자재명", "부자재", "자재명", "자재"}
ING_COLS = {
    "name": ING_NAME, "origin": {"원산지"}, "supplier": {"공급처", "업체", "업체명", "거래처"},
    "loss_moisture": {"수분loss", "수분로스", "수분"},
    "loss_split": {"소분loss", "소분로스", "소분"},
    "loss_sorting": {"선별loss", "선별로스", "선별"},
    "price": {"단가원kg", "단가", "벌크단가", "kg단가", "가격", "단가원"},
    "date": {"적용일", "기준일", "단가기준일", "날짜", "단가적용일"},
    "memo": {"메모", "비고"},
    "roasting": {"로스팅비", "로스팅단가", "로스팅비원kg", "로스팅단가원kg"},
    "sorting_cost": {"선별비", "선별단가", "선별비원kg", "선별단가원kg"},
    "currency": {"통화", "통화코드"},
    "foreign_price": {"외화단가", "외화단가kg", "외화단가원kg", "외화가격"},
    "fx_fixed": {"고정환율", "적용환율", "계약환율"},
}
MAT_COLS = {
    "name": MAT_NAME, "category": {"분류", "구분"}, "unit_price": {"단가원", "단가", "가격"},
    "loss_rate": {"loss", "로스", "loss율", "로스율", "기본loss"}, "memo": {"메모", "비고"},
    "date": {"적용일", "기준일", "단가기준일", "날짜", "단가적용일"},
    "currency": {"통화", "통화코드"}, "foreign_price": {"외화단가", "외화가격"},
    "fx_fixed": {"고정환율", "적용환율", "계약환율"},
}
FX_RATE_NAMES = {"환율", "환율원", "매매기준율", "기준환율"}
FX_COLS = {
    "currency": {"통화", "통화코드", "국가", "통화명"}, "date": {"기준일", "날짜", "적용일", "일자"},
    "rate": FX_RATE_NAMES, "memo": {"메모", "비고"},
}

# 앞 9개(기존 순서) + 새 항목은 뒤에 붙인다
ING_HEADERS = ["원료명", "원산지", "공급처", "수분 loss(%)", "소분 loss(%)", "선별 loss(%)",
               "단가(원/kg)", "적용일", "메모",
               "로스팅비(원/kg)", "선별비(원/kg)", "통화", "외화 단가(/kg)", "고정환율"]
MAT_HEADERS = ["부자재명", "분류", "단가(원)", "loss(%)", "메모",
               "적용일", "통화", "외화 단가", "고정환율"]
FX_HEADERS = ["통화", "기준일", "환율(원)", "메모"]


@dataclass
class ImportData:
    ingredients: list[dict] = field(default_factory=list)
    materials: list[dict] = field(default_factory=list)
    fx_rates: list[dict] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)


# ---------------------------------------------------------------- 파일 읽기
def _clean(v):
    if isinstance(v, str):
        v = v.replace(" ", " ").strip()
        return v or None
    return v


def _read_xlsx(data: bytes) -> dict[str, list[list]]:
    wb = load_workbook(BytesIO(data), data_only=True)
    out = {}
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows():
            r = []
            for c in row:
                v = c.value
                if isinstance(v, (int, float)) and not isinstance(v, bool) and "%" in (c.number_format or ""):
                    v = Pct(v * 100)
                r.append(_clean(v))
            rows.append(r)
        out[ws.title] = rows
    return out


def _read_xls(data: bytes) -> dict[str, list[list]]:
    try:
        import xlrd
    except ImportError as e:  # pragma: no cover
        raise RuntimeError(".xls 파일을 읽으려면 'pip install xlrd' 가 필요합니다 "
                           "(또는 .xlsx 로 저장해서 올려주세요)") from e
    book = xlrd.open_workbook(file_contents=data, formatting_info=True)
    out = {}
    for sh in book.sheets():
        rows = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    row.append(None)
                elif cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                elif cell.ctype == xlrd.XL_CELL_NUMBER:
                    fmt = book.format_map[book.xf_list[cell.xf_index].format_key].format_str
                    row.append(Pct(cell.value * 100) if "%" in fmt else cell.value)
                elif cell.ctype == xlrd.XL_CELL_ERROR:
                    row.append(None)
                else:
                    row.append(_clean(cell.value))
            rows.append(row)
        out[sh.name] = rows
    return out


def _read_csv(data: bytes, filename: str) -> dict[str, list[list]]:
    for enc in ("utf-8-sig", "cp949"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("CSV 문자 인코딩을 읽을 수 없습니다 (UTF-8 또는 CP949로 저장해 주세요)")
    rows = [[_clean(c) for c in r] for r in csv.reader(io.StringIO(text))]
    return {filename.rsplit(".", 1)[0]: rows}


def read_sheets(data: bytes, filename: str) -> dict[str, list[list]]:
    ext = filename.lower().rsplit(".", 1)[-1]
    if ext == "xlsx" or ext == "xlsm":
        return _read_xlsx(data)
    if ext == "xls":
        return _read_xls(data)
    if ext == "csv":
        return _read_csv(data, filename)
    raise ValueError("지원 형식: .xlsx / .xls / .csv")


# ---------------------------------------------------------------- 값 변환
def _num(v, what: str) -> float | None:
    if v is None:
        return None
    if isinstance(v, bool):
        raise ValueError(f"{what}: 숫자가 아닙니다")
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[,\s원%]", "", str(v))
    if s in ("", "-"):
        return None
    try:
        return float(s)
    except ValueError:
        raise ValueError(f"{what}: 숫자가 아닙니다 ('{v}')") from None


def _date(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    m = re.fullmatch(r"\s*(\d{4})[-./\s]+(\d{1,2})[-./\s]+(\d{1,2})\.?\s*", str(v))
    if not m:
        raise ValueError(f"적용일: 날짜 형식이 아닙니다 ('{v}', 예: 2026-09-30)")
    try:
        return date(int(m[1]), int(m[2]), int(m[3])).isoformat()
    except ValueError:
        raise ValueError(f"적용일: 존재하지 않는 날짜입니다 ('{v}')") from None


def _loss(v, what: str, raw_collector: list) -> float | None:
    """% 단위 입력 → 비율(0.05). 5 또는 '5%' 모두 5%."""
    n = _num(v, what)
    if n is None:
        return None
    if n < 0 or n > 100:
        raise ValueError(f"{what}: 0~100 사이여야 합니다 ({n})")
    if not isinstance(v, Pct):
        raw_collector.append(n)
    return n / 100


def _cur(v) -> str | None:
    s = _str(v)
    if s is None:
        return None
    key = re.sub(r"[\s()]", "", s).lower()
    if key in CUR_ALIASES:
        return CUR_ALIASES[key]
    code = re.sub(r"[^A-Za-z0-9]", "", s).upper()
    if 2 <= len(code) <= 6 and code.isalnum():
        return code
    raise ValueError(f"통화: 알 수 없는 값입니다 ('{s}', 예: USD, VND, TRY 또는 미국/베트남/튀르키예)")


def _pos(v, what: str) -> float | None:
    n = _num(v, what)
    if n is not None and n < 0:
        raise ValueError(f"{what}: 음수입니다")
    return n


def _str(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip() or None


# ---------------------------------------------------------------- 파싱
def _find_header(rows: list[list], name_set: set[str]) -> int | None:
    for i, row in enumerate(rows[:15]):
        if any(_norm(c) in name_set for c in row if c is not None):
            return i
    return None


def _map_columns(header: list, colspec: dict[str, set[str]]) -> tuple[dict[str, int], list[str]]:
    mapping, ignored = {}, []
    for idx, h in enumerate(header):
        if h is None:
            continue
        key = next((k for k, syn in colspec.items() if _norm(h) in syn), None)
        if key and key not in mapping:
            mapping[key] = idx
        else:
            ignored.append(str(h))
    return mapping, ignored


def parse_file(data: bytes, filename: str) -> ImportData:
    res = ImportData()
    try:
        sheets = read_sheets(data, filename)
    except Exception as e:
        res.errors.append(f"파일을 읽을 수 없습니다: {e}")
        return res

    ing_loss_raw: list[float] = []
    mat_loss_raw: list[float] = []

    for sname, rows in sheets.items():
        h_ing, h_mat = _find_header(rows, ING_NAME), _find_header(rows, MAT_NAME)
        h_fx = _find_header(rows, FX_RATE_NAMES) if (h_ing is None and h_mat is None) else None
        if h_ing is None and h_mat is None and h_fx is None:
            res.notes.append(f"시트 '{sname}': 헤더('원료명' / '부자재명' / '환율')가 없어 건너뜀")
            continue
        for kind, h, colspec in (("원료", h_ing, ING_COLS), ("부자재", h_mat, MAT_COLS), ("환율", h_fx, FX_COLS)):
            if h is None:
                continue
            mapping, ignored = _map_columns(rows[h], colspec)
            res.notes.append(f"시트 '{sname}' → {kind} 목록으로 인식 (헤더 {h + 1}행)")
            if ignored:
                res.notes.append(f"  └ 무시한 열: {', '.join(ignored)}")
            for i in range(h + 1, len(rows)):
                row = rows[i]
                if all(c is None for c in row):
                    continue
                get = lambda k: row[mapping[k]] if k in mapping and mapping[k] < len(row) else None
                src = f"{sname} {i + 1}행"
                if kind == "환율":
                    try:
                        cur, d, rate = _cur(get("currency")), _date(get("date")), _num(get("rate"), "환율")
                        if cur is None:
                            raise ValueError("통화가 비어 있습니다")
                        if d is None:
                            raise ValueError("기준일이 비어 있습니다")
                        if rate is None or rate <= 0:
                            raise ValueError("환율은 0보다 커야 합니다")
                        res.fx_rates.append({"src": src, "currency": cur, "date": d, "rate": rate,
                                             "memo": _str(get("memo"))})
                    except ValueError as e:
                        res.errors.append(f"{src}: {e}")
                    continue
                name = _str(get("name"))
                if not name:
                    res.errors.append(f"{src}: {kind}명이 비어 있습니다")
                    continue
                try:
                    cur, foreign = _cur(get("currency")), _pos(get("foreign_price"), "외화 단가")
                    fx_fixed = _pos(get("fx_fixed"), "고정환율")
                    if foreign is not None and cur in (None, "KRW"):
                        raise ValueError("외화 단가가 있는데 통화(USD/VND/TRY 등)가 없습니다")
                    if cur not in (None, "KRW") and foreign is None:
                        raise ValueError(f"통화가 {cur} 인데 외화 단가가 없습니다")
                    if fx_fixed is not None and fx_fixed <= 0:
                        raise ValueError("고정환율은 0보다 커야 합니다")
                    common = {"currency": cur, "foreign_price": foreign,
                              "fx_fixed": fx_fixed if foreign is not None else None}
                    if kind == "원료":
                        item = {
                            "name": name, "src": src,
                            "origin": _str(get("origin")), "supplier": _str(get("supplier")),
                            "loss_moisture": _loss(get("loss_moisture"), "수분 loss", ing_loss_raw),
                            "loss_split": _loss(get("loss_split"), "소분 loss", ing_loss_raw),
                            "loss_sorting": _loss(get("loss_sorting"), "선별 loss", ing_loss_raw),
                            "roasting_cost_per_kg": _pos(get("roasting"), "로스팅비"),
                            "sorting_cost_per_kg": _pos(get("sorting_cost"), "선별비"),
                            "price": _pos(get("price"), "단가"), "date": _date(get("date")),
                            "memo": _str(get("memo")), **common,
                        }
                        if item["price"] == 0 and foreign is None:
                            res.warnings.append(f"{src}: '{name}' 단가가 0원입니다")
                        if item["date"] and item["price"] is None and foreign is None:
                            res.warnings.append(f"{src}: '{name}' 적용일은 있는데 단가가 없어 무시됩니다")
                        for f in ("loss_moisture", "loss_split", "loss_sorting"):
                            if item[f] is not None and item[f] > 0.3:
                                res.warnings.append(f"{src}: '{name}' {f} 가 {item[f]*100:.1f}% 로 큽니다")
                        res.ingredients.append(item)
                    else:
                        cat_raw = _str(get("category"))
                        cat = None
                        if cat_raw:
                            cat = CAT_ALIASES.get(_norm(cat_raw))
                            if cat is None:
                                raise ValueError(f"분류: '{cat_raw}' → 롤포장지/포장지/인케이스/카톤박스/소분비/기타 중 하나여야 합니다")
                        item = {
                            "name": name, "src": src, "category": cat,
                            "unit_price": _pos(get("unit_price"), "단가"), "date": _date(get("date")),
                            "loss_rate": _loss(get("loss_rate"), "loss", mat_loss_raw),
                            "memo": _str(get("memo")), **common,
                        }
                        res.materials.append(item)
                except ValueError as e:
                    res.errors.append(f"{src} ({name}): {e}")

    for label, raw in (("원료", ing_loss_raw), ("부자재", mat_loss_raw)):
        pos = [x for x in raw if x > 0]
        if len(pos) >= 3 and all(x <= 1 for x in pos):
            res.warnings.append(
                f"{label} loss 값이 전부 1 이하입니다. 5%를 '0.05'로 적으셨다면 지금은 0.05% 로 읽힙니다. "
                f"'5' 또는 '5%'로 입력해야 5%예요. 미리보기에서 확인하세요.")
    if not res.ingredients and not res.materials and not res.fx_rates and not res.errors:
        res.errors.append("등록할 원료/부자재/환율 데이터를 찾지 못했습니다. 양식(템플릿)의 헤더를 확인하세요.")

    names = [r["name"] for r in res.materials]
    for n in sorted({n for n in names if names.count(n) > 1}):
        res.warnings.append(f"부자재 '{n}' 이(가) 파일 안에 여러 번 있습니다 (마지막 값이 적용됨)")
    keys = [(r["currency"], r["date"]) for r in res.fx_rates]
    for k in sorted({k for k in keys if keys.count(k) > 1}):
        res.warnings.append(f"환율 {k[0]} {k[1]} 이(가) 파일 안에 여러 번 있습니다 (마지막 값이 적용됨)")
    return res


# ---------------------------------------------------------------- 양식 / 내보내기
_HEAD = PatternFill("solid", fgColor="DDEBF7")
_NOTE = PatternFill("solid", fgColor="FFF2CC")


def _header(ws, headers):
    ws.append(headers)
    for i, c in enumerate(ws[1], 1):
        c.font, c.fill = Font(bold=True), _HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(headers[i - 1]) * 2.2)
    ws.freeze_panes = "A2"


def build_workbook(ingredients: list[dict] | None = None, materials: list[dict] | None = None,
                   fx_rates: list[dict] | None = None) -> bytes:
    """빈 양식(인자 없음) 또는 현재 마스터를 채운 파일. 시트: 작성방법 / 원료 / 부자재 / 환율."""
    wb = Workbook()
    guide = wb.active
    guide.title = "작성방법"
    for line in [
        "■ 원료·부자재·환율 일괄 등록 양식",
        "",
        "1. '원료', '부자재', '환율' 시트에 한 줄에 하나씩 입력합니다. (헤더 이름은 바꾸지 마세요)",
        "2. 이름이 같은 원료/부자재가 이미 있으면 '수정', 없으면 '신규 등록'됩니다. 환율은 (통화+기준일)이 같으면 수정.",
        "3. 빈 칸은 '기존 값 유지'입니다. (신규일 때만 0 또는 빈 값으로 등록)",
        "4. loss(%)는 퍼센트 숫자로 입력: 5% → 5  (엑셀 %서식 셀도 인식합니다)",
        "5. 단가를 입력하면 '단가 이력'에 적용일과 함께 추가됩니다. 적용일을 비우면 오늘 날짜입니다.",
        "   같은 품목의 단가를 날짜별로 여러 줄 넣어도 됩니다 (가격 추이·변동률 조회에 사용됩니다).",
        "6. 수입 품목은 '통화'(USD/VND/TRY 또는 미국/베트남/튀르키예, 그 외는 환율 메뉴에 추가한 코드)와",
        "   '외화 단가'를 입력하세요. 원화 단가 칸은 비워도 됩니다(외화 단가가 있으면 무시).",
        "   환율은 그 날짜의 '환율' 시트/환율 메뉴 값으로 자동 환산됩니다. 계약환율을 쓰려면 '고정환율'(원)을 입력하세요.",
        "7. 로스팅비(원/kg), 선별비(원/kg)는 원료별 가공비입니다 (없으면 비워 두세요).",
        "8. 부자재 '분류': 롤포장지 / 포장지 / 인케이스 / 카톤박스 / 소분비 / 기타 (신규 부자재는 필수)",
        "9. 환율은 '원 / 외화 1단위'로 입력합니다. 단, 베트남 동(VND)은 100동당 원으로 입력합니다 (예: 5.5).",
        "10. 오류가 한 줄이라도 있으면 전체 등록이 취소되고, 어느 시트 몇 행인지 알려줍니다.",
        "11. 등록 전에 미리보기로 신규/수정/변경없음을 확인할 수 있습니다.",
        "",
        "예) 원료(원화):  아몬드 | 미국 | OO상사 | 2.5 | 3 | 0.54 | 10500 | 2026-09-30 | 9월 견적",
        "예) 원료(수입):  아몬드 | … | 단가(원/kg) 비움 | 적용일 2026-09-30 | … | 통화 USD | 외화 단가 7.5",
        "예) 부자재:      소포장지 | 소포장지 | 20 | 3 |",
        "예) 환율:        USD | 2026-09-30 | 1385.5",
    ]:
        guide.append([line])
    guide["A1"].font = Font(bold=True, size=13)
    guide.column_dimensions["A"].width = 110

    wi = wb.create_sheet("원료")
    _header(wi, ING_HEADERS)
    for r in ingredients or []:
        foreign = r.get("currency") not in (None, "KRW") and r.get("foreign_price") is not None
        wi.append([r["name"], r.get("origin") or "", r.get("supplier") or "",
                   round(r.get("loss_moisture", 0) * 100, 4), round(r.get("loss_split", 0) * 100, 4),
                   round(r.get("loss_sorting", 0) * 100, 4),
                   None if foreign else r.get("price_per_kg"),
                   r.get("price_date") or "", "",
                   r.get("roasting_cost_per_kg") or None, r.get("sorting_cost_per_kg") or None,
                   r.get("currency") if foreign else None, r.get("foreign_price") if foreign else None,
                   r.get("fx_fixed_rate") if foreign and r.get("fx_mode") == "fixed" else None])

    wm = wb.create_sheet("부자재")
    _header(wm, MAT_HEADERS)
    for r in materials or []:
        foreign = r.get("currency") not in (None, "KRW") and r.get("foreign_price") is not None
        wm.append([r["name"], CATEGORIES[r["category"]], None if foreign else r["unit_price"],
                   round(r["loss_rate"] * 100, 4), r.get("memo") or "",
                   r.get("price_date") or "", r.get("currency") if foreign else None,
                   r.get("foreign_price") if foreign else None,
                   r.get("fx_fixed_rate") if foreign and r.get("fx_mode") == "fixed" else None])
    dv = DataValidation(type="list", formula1='"' + ",".join(CATEGORIES.values()) + '"', allow_blank=True)
    wm.add_data_validation(dv)
    dv.add("B2:B2000")

    wf = wb.create_sheet("환율")
    _header(wf, FX_HEADERS)
    for r in fx_rates or []:
        wf.append([r["currency"], r["rate_date"], r["rate"], r.get("memo") or ""])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
