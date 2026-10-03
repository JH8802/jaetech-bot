"""견적 결과를 엑셀(.xlsx)로 내보내기."""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .engine import CATEGORIES, COST_LINES, QuoteInput, QuoteResult, calc_mode

HEAD = PatternFill("solid", fgColor="DDEBF7")
BOLD = Font(bold=True)
MONEY = '#,##0'
MONEY1 = '#,##0.0'
PCT = '0.0%'
PCT2 = '0.00%'


def _fit(ws, maxw=40):
    for col in ws.columns:
        w = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(10, w * 1.6), maxw)


def _head(ws, row=1):
    for c in ws[row]:
        c.font, c.fill = BOLD, HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _save(wb) -> bytes:
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _norm_items(items):
    """(q, r) 또는 (q, r, meta) 를 (q, r, meta|None) 으로 통일."""
    return [(i[0], i[1], i[2] if len(i) > 2 else None) for i in items]


def _fmt_money_cols(ws, cols, fmt=MONEY1):
    for col in cols:
        ws.cell(ws.max_row, col).number_format = fmt


# =====================================================================
# 견적 1건(또는 여러 건) 원가표
# =====================================================================
def _lines(q, r, only_nonzero_optional: bool = True):
    """(라벨, 값) 목록. 선택 항목(로스팅·3PL·수수료 등)은 0이면 생략."""
    out = []
    for label, attr, core in COST_LINES:
        v = getattr(r, attr)
        if core or not only_nonzero_optional or abs(v) > 1e-9:
            if attr == "sga":
                label += f" ({q.sga_rate:.2%})"
            elif attr == "margin":
                label += f" ({q.margin_rate:.2%})"
            elif attr == "logistics":
                label += f" ({q.logistics_rate:.2%}{' VAT포함기준' if q.logistics_vat else ''}" \
                         + (f" + 정액 {q.logistics_fixed:,.0f}" if q.logistics_fixed else "") + ")"
            elif attr == "fee":
                label += f" ({q.fee_rate:.2%}{' VAT포함기준' if q.online_vat else ''})"
            elif attr == "promo":
                label += f" ({q.promo_rate:.2%})"
            elif attr == "event":
                label += f" ({q.event_rate:.2%}" + (f" + 정액 {q.event_fixed:,.0f}" if q.event_fixed else "") + ")"
            out.append((label, v))
    return out


def quotes_to_xlsx(items, internal: bool = True) -> bytes:
    """internal=True : 원가 내역 포함(사내용)
       internal=False: 제품명/납품가만 (거래처 제출용)"""
    items = _norm_items(items)
    wb = Workbook()
    if not internal:
        ws = wb.active
        ws.title = "납품가"
        ws.append(["제품명", "판매처", "납품가(원, VAT 별도)"])
        _head(ws)
        for q, r, _m in items:
            ws.append([q.product_name, q.channel_name or "", r.price])
            ws.cell(ws.max_row, 3).number_format = MONEY
        _fit(ws)
        return _save(wb)

    # ---- 요약 시트 ----
    ws = wb.active
    ws.title = "요약"
    heads = ["제품명", "판매처", "계산방식", "견적기준일"] + [l for l, _a, _c in COST_LINES] + ["실질마진율"]
    ws.append(heads)
    _head(ws)
    for q, r, _m in items:
        ws.append([q.product_name, q.channel_name or "(기본 설정)", calc_mode(q), q.as_of_date]
                  + [getattr(r, a) for _l, a, _c in COST_LINES] + [r.effective_margin_rate])
        _fmt_money_cols(ws, range(5, 5 + len(COST_LINES)))
        ws.cell(ws.max_row, len(heads)).number_format = PCT
    _fit(ws)

    # ---- 제품별 상세 시트 ----
    used = set()
    for q, r, meta in items:
        title = q.product_name[:28]
        for ch in '/\\*?[]:':
            title = title.replace(ch, "_")
        base, n = title, 2
        while title in used:
            title, n = f"{base[:25]}_{n}", n + 1
        used.add(title)
        d = wb.create_sheet(title)
        d.append(["제품명", q.product_name, "", "봉수", q.bag_count, "1봉 중량(g)", q.unit_weight_g])
        info = ["계산방식", calc_mode(q), "", "판매처", q.channel_name or "(기본 설정)",
                "견적 기준일", q.as_of_date or "-"]
        d.append(info)
        if meta:
            d.append(["견적번호", meta.get("id"), "", "작성일시", meta.get("created_at"),
                      "작성자", meta.get("created_by"), "메모", meta.get("memo")])
        d.append([])
        d.append(["원료", "구성비", "단위중량(g)", "벌크단가(원/kg)", "원물가", "로스율", "로스", "로스팅비", "선별비",
                  "원산지", "공급처", "단가기준일", "단가출처", "통화", "외화단가", "적용환율(원/1단위)", "환율기준일"])
        _head(d, d.max_row)
        for x in r.ingredient_rows:
            d.append([x["name"], x["ratio"], x["unit_g"], x["price_per_kg"], x["cost"],
                      x["loss_rate"], x["loss"], x.get("roasting", 0.0), x.get("sorting", 0.0),
                      x.get("origin", ""), x.get("supplier", ""), x.get("price_date", ""),
                      x.get("price_source", ""), x.get("currency", "KRW"), x.get("foreign_price"),
                      x.get("fx_rate"), x.get("fx_date", "")])
            d.cell(d.max_row, 2).number_format = PCT
            d.cell(d.max_row, 6).number_format = PCT
            _fmt_money_cols(d, (4, 5, 7, 8, 9))
            d.cell(d.max_row, 15).number_format = '#,##0.00'
            d.cell(d.max_row, 16).number_format = '#,##0.0000'
        d.append([])
        d.append(["부자재", "분류", "수량", "봉수연동", "기본금액", "loss", "합계", "단가기준일", "통화", "외화단가",
                  "적용환율(원/1단위)"])
        _head(d, d.max_row)
        for x in r.material_rows:
            d.append([x["name"], CATEGORIES[x["category"]], x.get("qty"), "예" if x.get("per_bag") else "",
                      x["base"], x["loss"], x["total"], x.get("price_date", ""), x.get("currency", "KRW"),
                      x.get("foreign_price"), x.get("fx_rate")])
            _fmt_money_cols(d, (5, 6, 7))
            d.cell(d.max_row, 10).number_format = '#,##0.00'
            d.cell(d.max_row, 11).number_format = '#,##0.0000'
        d.append([])
        for label, val in _lines(q, r):
            d.append([label, val])
            d.cell(d.max_row, 1).font = BOLD
            d.cell(d.max_row, 2).number_format = MONEY1
        d.append(["(참고) 부자재 loss 합계 — 각 부자재 금액에 포함", r.material_loss_cost])
        d.cell(d.max_row, 2).number_format = MONEY1
        d.append(["납품가 반올림", q.rounding])
        for w in r.warnings:
            d.append(["⚠ " + w])
        _fit(d)
    return _save(wb)


# =====================================================================
# 견적 이력 목록 전체 (조건 검색 결과)
# =====================================================================
def quotes_list_to_xlsx(items) -> bytes:
    """items: db.search_quotes() 결과 [(q, r, meta), ...].  사내용(원가 포함)."""
    items = _norm_items(items)
    wb = Workbook()
    ws = wb.active
    ws.title = "견적목록"
    heads = (["견적번호", "일시", "작성자", "제품", "계산방식", "판매처", "견적기준일", "봉수", "1봉중량(g)"]
             + [l for l, _a, _c in COST_LINES]
             + ["판관비율", "마진율", "물류비율", "물류비 VAT기준", "수수료율", "홍보비율", "기획전비율",
                "반올림", "실질마진율", "경고", "메모"])
    ws.append(heads)
    _head(ws)
    wi = wb.create_sheet("원료상세")
    wi.append(["견적번호", "일시", "제품", "원료", "원산지", "공급처", "구성비", "벌크단가(원/kg)",
               "단가기준일", "단가출처", "통화", "외화단가", "적용환율(원/1단위)", "원물가", "로스율", "로스",
               "로스팅비", "선별비"])
    _head(wi)
    wm = wb.create_sheet("부자재상세")
    wm.append(["견적번호", "일시", "제품", "부자재", "분류", "수량", "봉수연동", "기본금액", "loss", "합계",
               "통화", "외화단가", "적용환율(원/1단위)"])
    _head(wm)

    n0 = 10
    for q, r, m in items:
        m = m or {}
        ws.append([m.get("id"), m.get("created_at"), m.get("created_by"), q.product_name, calc_mode(q),
                   q.channel_name or "(기본 설정)", q.as_of_date, q.bag_count, q.unit_weight_g]
                  + [getattr(r, a) for _l, a, _c in COST_LINES]
                  + [q.sga_rate, q.margin_rate, q.logistics_rate, "예" if q.logistics_vat else "아니오",
                     q.fee_rate, q.promo_rate, q.event_rate, q.rounding, r.effective_margin_rate,
                     " / ".join(r.warnings), m.get("memo")])
        row = ws.max_row
        _fmt_money_cols(ws, range(n0, n0 + len(COST_LINES)))
        base = n0 + len(COST_LINES)
        for col in (base, base + 1, base + 2, base + 4, base + 5, base + 6):
            ws.cell(row, col).number_format = PCT2
        ws.cell(row, base + 8).number_format = PCT

        for x in r.ingredient_rows:
            wi.append([m.get("id"), m.get("created_at"), q.product_name, x["name"],
                       x.get("origin", ""), x.get("supplier", ""), x["ratio"], x["price_per_kg"],
                       x.get("price_date", ""), x.get("price_source", ""), x.get("currency", "KRW"),
                       x.get("foreign_price"), x.get("fx_rate"), x["cost"], x["loss_rate"], x["loss"],
                       x.get("roasting", 0.0), x.get("sorting", 0.0)])
            wi.cell(wi.max_row, 7).number_format = PCT
            wi.cell(wi.max_row, 15).number_format = PCT
            _fmt_money_cols(wi, (8, 14, 16, 17, 18))
            wi.cell(wi.max_row, 12).number_format = '#,##0.00'
            wi.cell(wi.max_row, 13).number_format = '#,##0.0000'
        for x in r.material_rows:
            wm.append([m.get("id"), m.get("created_at"), q.product_name, x["name"],
                       CATEGORIES[x["category"]], x.get("qty"), "예" if x.get("per_bag") else "",
                       x["base"], x["loss"], x["total"], x.get("currency", "KRW"),
                       x.get("foreign_price"), x.get("fx_rate")])
            _fmt_money_cols(wm, (8, 9, 10))

    for sheet in (ws, wi, wm):
        sheet.freeze_panes = "A2"
        _fit(sheet, 30)
    ws.auto_filter.ref = ws.dimensions
    wi.auto_filter.ref = wi.dimensions
    wm.auto_filter.ref = wm.dimensions
    return _save(wb)


# =====================================================================
# 견적 2건 비교
# =====================================================================
def compare_to_xlsx(cmp: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "요약"
    ma, mb = cmp["meta_a"], cmp["meta_b"]
    ws.append([cmp["headline"]])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([f"A(기준): #{ma.get('id')} {ma.get('product_name')} {ma.get('created_at')} {ma.get('created_by')} [{ma.get('mode')}]"])
    ws.append([f"B(비교): #{mb.get('id')} {mb.get('product_name')} {mb.get('created_at')} {mb.get('created_by')} [{mb.get('mode')}]"])
    ws.append([])
    ws.append(["항목", "A", "B", "차이", "차이%"])
    _head(ws, ws.max_row)
    for r in cmp["summary"]:
        ws.append([r["항목"], r["A"], r["B"], r["차이"], r["차이%"]])
        is_rate = r["항목"] == "실질 마진율"
        for col in (2, 3, 4):
            ws.cell(ws.max_row, col).number_format = PCT if is_rate else MONEY1
        ws.cell(ws.max_row, 5).number_format = PCT
    _fit(ws, 60)

    def table(title, rows, money_cols=(), pct_cols=()):
        sh = wb.create_sheet(title)
        if not rows:
            sh.append(["(내용 없음)"])
            return
        keys = list(rows[0])
        sh.append(keys)
        _head(sh)
        for row in rows:
            sh.append([row[k] for k in keys])
            for i, k in enumerate(keys, 1):
                if k in money_cols:
                    sh.cell(sh.max_row, i).number_format = MONEY1
                if k in pct_cols:
                    sh.cell(sh.max_row, i).number_format = PCT
        _fit(sh)

    table("원료 비교", cmp["ingredients"],
          money_cols={"단가 A", "단가 B", "외화단가 A", "외화단가 B", "환율 A", "환율 B", "원료비 A", "원료비 B", "차이"},
          pct_cols={"구성비 A", "구성비 B"})
    table("부자재 비교", cmp["materials"], money_cols={"A", "B", "차이"})
    table("적용 규칙", cmp["rules"])
    table("변동 요인", [{"순위": i + 1, "내용": t} for i, t in enumerate(cmp["drivers"])])
    return _save(wb)


# =====================================================================
# 가격 추이 / 환율 / 제품 원가 추이 엑셀 (표 + 엑셀 그래프)
# =====================================================================
def table_workbook(sheets: list[dict]) -> bytes:
    """sheets: [{title, headers, rows, fmts={헤더: 서식}, widths?, chart?}]
       chart = {kind:'line'|'bar', x:'헤더', ys:['헤더',...], title, y_title}  (엑셀 그래프 삽입)"""
    from openpyxl.chart import BarChart, LineChart, Reference

    wb = Workbook()
    wb.remove(wb.active)
    for spec in sheets:
        title = spec["title"][:31]
        for ch in '/\\*?[]:':
            title = title.replace(ch, "_")
        ws = wb.create_sheet(title)
        headers, rows, fmts = spec["headers"], spec["rows"], spec.get("fmts", {})
        ws.append(headers)
        _head(ws)
        for row in rows:
            ws.append(list(row))
        for ci, h in enumerate(headers, 1):
            if h in fmts:
                for r in range(2, ws.max_row + 1):
                    ws.cell(r, ci).number_format = fmts[h]
        ws.freeze_panes = "A2"
        _fit(ws, 45)
        chart = spec.get("chart")
        if chart and rows:
            ch = LineChart() if chart["kind"] == "line" else BarChart()
            ch.title = chart.get("title", "")
            ch.y_axis.title = chart.get("y_title", "")
            ch.height, ch.width = 9, 22
            xcol = headers.index(chart["x"]) + 1
            for yh in chart["ys"]:
                ycol = headers.index(yh) + 1
                ch.add_data(Reference(ws, min_col=ycol, min_row=1, max_row=ws.max_row), titles_from_data=True)
            ch.set_categories(Reference(ws, min_col=xcol, min_row=2, max_row=ws.max_row))
            ws.add_chart(ch, f"{get_column_letter(len(headers) + 2)}2")
        if not rows:
            ws.append(["(조회 결과 없음)"])
    return _save(wb)


def item_trend_xlsx(name: str, kind_label: str, series, stats, dec, agg, freq_label: str) -> bytes:
    sheets = [
        {"title": "일별 가격", "headers": ["날짜", "원화 단가", "통화", "외화 단가", "적용 환율(원/1단위)", "환율 기준일",
                                         "단가 적용일", "변동일"],
         "rows": [[r["date"], r["krw"], r["currency"], r["foreign"], r["fx"], r["fx_date"], r["price_date"],
                   "●" if r["changed"] else ""] for r in series],
         "fmts": {"원화 단가": MONEY1, "외화 단가": '#,##0.00', "적용 환율(원/1단위)": '#,##0.0000'},
         "chart": {"kind": "line", "x": "날짜", "ys": ["원화 단가"], "title": f"{name} 원화 단가 추이", "y_title": "원"}},
        {"title": f"{freq_label} 집계", "headers": ["기간", "평균", "최저", "최고", "기말"],
         "rows": [[a["period"], a["avg"], a["min"], a["max"], a["last"]] for a in agg],
         "fmts": {"평균": MONEY1, "최저": MONEY1, "최고": MONEY1, "기말": MONEY1},
         "chart": {"kind": "bar", "x": "기간", "ys": ["평균"], "title": f"{name} {freq_label} 평균 단가", "y_title": "원"}},
    ]
    if stats:
        rows = [["품목", f"{name} ({kind_label})"], ["기간", f"{stats['start_date']} ~ {stats['end_date']}"],
                ["시작 단가(원)", stats["start"]], ["종료 단가(원)", stats["end"]],
                ["변동액(원)", stats["delta"]], ["변동률", stats["rate"]],
                ["최저(원)", stats["min"]], ["최저일", stats["min_date"]],
                ["최고(원)", stats["max"]], ["최고일", stats["max_date"]],
                ["평균(원)", stats["avg"]], ["가격이 바뀐 날 수", stats["change_days"]]]
        if dec:
            rows += [["", ""], ["[환율 영향 분해]", dec["currency"]],
                     ["현지 단가 변동률", dec["foreign_rate"]], ["환율 변동률", dec["fx_rate"]],
                     ["원화 단가 변동률", dec["krw_rate"]],
                     ["현지가 변동 효과(원)", dec["price_effect"]], ["환율 변동 효과(원)", dec["fx_effect"]],
                     ["합계(원)", dec["total"]]]
        sheets.insert(0, {"title": "요약", "headers": ["항목", "값"], "rows": rows})
    if dec:
        sheets.append({"title": "지수(기준=100)", "headers": ["date", "원화 단가", "현지 단가", "환율"],
                       "rows": [[r["date"], r["원화 단가"], r["현지 단가"], r["환율"]] for r in dec["index"]],
                       "fmts": {"원화 단가": '0.0', "현지 단가": '0.0', "환율": '0.0'},
                       "chart": {"kind": "line", "x": "date", "ys": ["원화 단가", "현지 단가", "환율"],
                                 "title": "시작일=100 지수 (원화 단가 = 현지 단가 x 환율)", "y_title": "지수"}})
    wb_bytes = table_workbook(sheets)
    wb = load_workbook(BytesIO(wb_bytes))
    if "요약" in wb.sheetnames:                         # 요약 시트 서식
        ws = wb["요약"]
        for r in range(2, ws.max_row + 1):
            label, v = ws.cell(r, 1).value or "", ws.cell(r, 2).value
            if isinstance(v, float):
                ws.cell(r, 2).number_format = PCT if "률" in label else MONEY1
    return _save(wb)


def ranking_xlsx(rows: list[dict], kind_label: str, d_from: str, d_to: str) -> bytes:
    return table_workbook([{
        "title": f"{kind_label} 변동률 순위",
        "headers": ["품목", "시작일", "종료일", "시작 단가", "종료 단가", "변동액", "변동률", "최저", "최고", "평균",
                    "가격 변동일 수", "통화", "현지가 변동률", "환율 변동률", "현지가 효과(원)", "환율 효과(원)"],
        "rows": [[r["name"], r["start_date"], r["end_date"], r["start"], r["end"], r["delta"], r["rate"],
                  r["min"], r["max"], r["avg"], r["change_days"], r["currency"], r["foreign_rate"],
                  r["fx_rate"], r["price_effect"], r["fx_effect"]] for r in rows],
        "fmts": {"시작 단가": MONEY1, "종료 단가": MONEY1, "변동액": MONEY1, "변동률": PCT, "최저": MONEY1,
                 "최고": MONEY1, "평균": MONEY1, "현지가 변동률": PCT, "환율 변동률": PCT,
                 "현지가 효과(원)": MONEY1, "환율 효과(원)": MONEY1},
        "chart": {"kind": "bar", "x": "품목", "ys": ["변동률"], "title": f"{d_from} ~ {d_to} 변동률", "y_title": "변동률"},
    }])


def yearly_xlsx(rows: list[dict], kind_label: str) -> bytes:
    long_rows = [[r["name"], r["year"], r["base_label"], r["base"], r["end"], r["delta"], r["rate"], r["avg"],
                  r["min"], r["max"], "진행 중(연도 미경과)" if r["partial"] else ""] for r in rows]
    years = sorted({r["year"] for r in rows})
    names = sorted({r["name"] for r in rows})
    lookup = {(r["name"], r["year"]): r["rate"] for r in rows}
    pivot = [[n] + [lookup.get((n, y)) for y in years] for n in names]
    return table_workbook([
        {"title": f"{kind_label} 연도별 변동", "headers": ["품목", "연도", "기준 시점", "기준 단가", "연말 단가", "변동액",
                                                          "변동률", "연평균", "연중 최저", "연중 최고", "비고"],
         "rows": long_rows,
         "fmts": {"기준 단가": MONEY1, "연말 단가": MONEY1, "변동액": MONEY1, "변동률": PCT, "연평균": MONEY1,
                  "연중 최저": MONEY1, "연중 최고": MONEY1}},
        {"title": "변동률 표", "headers": ["품목"] + [f"{y}년" for y in years], "rows": pivot,
         "fmts": {f"{y}년": PCT for y in years}},
    ])


def fx_xlsx(per_currency: dict[str, list[dict]], names: dict[str, str]) -> bytes:
    sheets = []
    for cur, series in per_currency.items():
        sheets.append({
            "title": f"{cur} 환율", "headers": ["날짜", "환율(원)", "환율 기준일", "변동일"],
            "rows": [[r["date"], r["rate"], r["rate_date"], "●" if r["changed"] else ""] for r in series],
            "fmts": {"환율(원)": '#,##0.0000'},
            "chart": {"kind": "line", "x": "날짜", "ys": ["환율(원)"],
                      "title": f"{names.get(cur, cur)} 환율", "y_title": "원"}})
    return table_workbook(sheets)


def product_trend_xlsx(name: str, rows: list[dict]) -> bytes:
    heads = ["날짜", "납품가", "원가합계", "직접원가", "원료(원물가+로스)", "로스팅·선별", "부자재", "운송·3PL",
             "판매처 비용", "실질 마진", "실질 마진율", "비고"]
    return table_workbook([{
        "title": "제품 원가 추이", "headers": heads,
        "rows": [[r["date"], r["price"], r["total_cost"], r["direct_cost"], r["ingredients"], r["roasting_sorting"],
                  r["materials"], r["logistics_etc"], r["channel_cost"], r["effective_margin"], r["margin_rate"],
                  r["notes"]] for r in rows],
        "fmts": {h: MONEY1 for h in heads[1:10]} | {"실질 마진율": PCT},
        "chart": {"kind": "line", "x": "날짜", "ys": ["납품가", "직접원가"], "title": f"{name} 납품가·직접원가 추이",
                  "y_title": "원"}}])


# =====================================================================
# 전체 데이터 엑셀 (백업 · 외부 분석용)
# =====================================================================
def all_data_xlsx(data: dict, only: list[str] | None = None) -> bytes:
    """db.collect_all_data() 결과를 시트별로 내보낸다. 원료·부자재 단가 이력, 환율 이력, 판매처, 제품(BOM), 견적 이력 포함."""
    cat = CATEGORIES
    kinds = {"offline": "오프라인(마트)", "online": "온라인"}
    pct_cols = {"로스율", "수분loss", "소분loss", "선별loss", "구성비", "loss", "물류비율", "수수료율", "홍보비율",
                "기획전비율", "판관비율", "마진율"}
    def pc(v):
        return None if v is None else v

    sheets = []
    sheets.append({"title": "원료 마스터", "headers": ["원료", "원산지", "공급처", "수분loss", "소분loss", "선별loss",
                                                       "로스팅비(원/kg)", "선별비(원/kg)", "현재단가(원/kg)", "통화",
                                                       "외화단가", "적용환율", "단가적용일"],
                   "rows": [[r["name"], r["origin"], r["supplier"], r["loss_moisture"], r["loss_split"],
                             r["loss_sorting"], r["roasting_cost_per_kg"], r["sorting_cost_per_kg"], r["price_per_kg"],
                             r["currency"], r["foreign_price"], r["fx"], r["price_date"]] for r in data["ingredients"]],
                   "fmts": {"수분loss": PCT2, "소분loss": PCT2, "선별loss": PCT2, "현재단가(원/kg)": MONEY1,
                            "외화단가": '#,##0.00', "적용환율": '#,##0.0000'}})
    sheets.append({"title": "원료 단가이력", "headers": ["원료", "적용일", "통화", "외화단가", "환율방식", "고정환율",
                                                       "원화단가(입력시 환산)", "메모"],
                   "rows": [[r["name"], r["effective_date"], r["currency"] or "KRW", r["foreign_price"],
                             ("고정" if r["fx_mode"] == "fixed" else "자동") if r["foreign_price"] is not None else "",
                             r["fx_fixed_rate"], r["price_per_kg"], r["memo"]] for r in data["ingredient_prices"]],
                   "fmts": {"외화단가": '#,##0.00', "고정환율": '#,##0.0000', "원화단가(입력시 환산)": MONEY1}})
    sheets.append({"title": "부자재 마스터", "headers": ["부자재", "분류", "현재단가(원)", "통화", "외화단가", "적용환율",
                                                        "단가적용일", "loss", "메모"],
                   "rows": [[r["name"], cat.get(r["category"], r["category"]), r["unit_price"], r["currency"],
                             r["foreign_price"], r["fx"], r["price_date"], r["loss_rate"], r["memo"]]
                            for r in data["materials"]],
                   "fmts": {"현재단가(원)": MONEY1, "외화단가": '#,##0.00', "적용환율": '#,##0.0000', "loss": PCT2}})
    sheets.append({"title": "부자재 단가이력", "headers": ["부자재", "분류", "적용일", "통화", "외화단가", "환율방식",
                                                         "고정환율", "원화단가(입력시 환산)", "메모"],
                   "rows": [[r["name"], cat.get(r["category"], r["category"]), r["effective_date"],
                             r["currency"] or "KRW", r["foreign_price"],
                             ("고정" if r["fx_mode"] == "fixed" else "자동") if r["foreign_price"] is not None else "",
                             r["fx_fixed_rate"], r["unit_price"], r["memo"]] for r in data["material_prices"]],
                   "fmts": {"외화단가": '#,##0.00', "고정환율": '#,##0.0000', "원화단가(입력시 환산)": MONEY1}})
    units = {c["code"]: c["quote_unit"] for c in data["currencies"]}
    sheets.append({"title": "환율 이력", "headers": ["통화", "기준일", "환율(원)", "표기단위", "메모"],
                   "rows": [[r["currency"], r["rate_date"], r["rate"],
                             f"외화 {units.get(r['currency'], 1):g}단위당 원", r["memo"]] for r in data["fx_rates"]],
                   "fmts": {"환율(원)": '#,##0.0000'}})
    sheets.append({"title": "통화", "headers": ["코드", "이름", "국가", "표기단위", "구분"],
                   "rows": [[c["code"], c["name"], c["country"], c["quote_unit"], "기본" if c["builtin"] else "(기타)"]
                            for c in data["currencies"]]})
    sheets.append({"title": "판매처", "headers": ["판매처", "유형", "물류비율", "정액물류비(원/개)", "물류비 VAT기준",
                                                 "수수료율", "홍보비율", "기획전비율", "기획전정액(원/개)", "택배비(원/개)",
                                                 "수수료류 VAT기준", "판관비율", "마진율", "메모"],
                   "rows": [[c["name"], kinds.get(c["kind"], c["kind"]), c["logistics_rate"], c["logistics_fixed"],
                             "예" if c["logistics_vat"] else "", c["fee_rate"], c["promo_rate"], c["event_rate"],
                             c["event_fixed"], c["parcel_cost"], "예" if c["online_vat"] else "",
                             c["sga_rate"], c["margin_rate"], c["memo"]] for c in data["channels"]],
                   "fmts": {k: PCT2 for k in ("물류비율", "수수료율", "홍보비율", "기획전비율", "판관비율", "마진율")}})
    sheets.append({"title": "제품", "headers": ["제품", "봉수", "1봉중량(g)", "공통 선별단가(원/kg)", "운송비(원)",
                                               "3PL 이용료(원)", "메모"],
                   "rows": [[p["name"], p["bag_count"], p["unit_weight_g"], p["sorting_cost_per_kg"],
                             p["shipping_cost"], p.get("three_pl_cost") or 0, p["memo"]] for p in data["products"]]})
    bom = [[r["product"], "원료", r["item"], "", r["ratio"], None, "", r["loss_rate"], r["price_override"]]
           for r in data["bom_ingredients"]]
    bom += [[r["product"], "부자재", r["item"], cat.get(r["category"], r["category"]), None, r["qty"],
             "봉수 연동" if r["qty_basis"] == "bag" else "고정", r["loss_rate"], None] for r in data["bom_materials"]]
    sheets.append({"title": "제품 구성(BOM)", "headers": ["제품", "구분", "항목", "분류", "구성비", "수량", "수량기준",
                                                         "loss/로스율(개별설정)", "단가 직접입력(원/kg)"],
                   "rows": sorted(bom, key=lambda r: (r[0], r[1] != "원료")),
                   "fmts": {"구성비": PCT, "loss/로스율(개별설정)": PCT2}})
    sheets.append({"title": "견적 이력 요약", "headers": ["견적번호", "일시", "작성자", "제품", "계산방식", "판매처",
                                                        "견적기준일", "원가합계", "납품가", "실질마진", "실질마진율", "메모"],
                   "rows": [[m["id"], m["created_at"], m["created_by"], q.product_name, calc_mode(q),
                             q.channel_name or "(기본 설정)", q.as_of_date, r.total_cost, r.price,
                             r.effective_margin, r.effective_margin_rate, m["memo"]]
                            for q, r, m in data.get("quotes", [])],
                   "fmts": {"원가합계": MONEY1, "납품가": MONEY1, "실질마진": MONEY1, "실질마진율": PCT}})
    if only:
        sheets = [sh for sh in sheets if sh["title"] in only]
    return table_workbook(sheets)
