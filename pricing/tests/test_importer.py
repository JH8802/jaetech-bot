"""엑셀 일괄 등록 테스트 (가짜 데이터)."""
import io
import pytest
from openpyxl import Workbook, load_workbook
from pricing import db, importer


@pytest.fixture()
def tmpdb(tmp_path, monkeypatch):
    monkeypatch.setattr(db, "DB_PATH", tmp_path / "t.db")
    db.init_db()


def make_xlsx(ing_rows=(), mat_rows=(), title_rows=0):
    wb = Workbook()
    wi = wb.active
    wi.title = "원료 단가표"                      # 시트 이름은 자유
    for _ in range(title_rows):
        wi.append(["2026년 9월 원료 단가표 (제목줄)"])
    wi.append(importer.ING_HEADERS)
    for r in ing_rows:
        wi.append(list(r))
    wm = wb.create_sheet("부자재목록")
    wm.append(importer.MAT_HEADERS)
    for r in mat_rows:
        wm.append(list(r))
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


ING = [("아몬드", "미국", "가상상사", 2.5, 3, 0.54, 10500, "2026-09-30", "9월"),
       ("호두", None, None, 0, 3, 0.5, "9,800원", None, None)]
MAT = [("소포장지", "포장지", 20, 3, None), ("카톤", "박스", 1000, "3%", "20입")]


def test_parse_basic():
    p = importer.parse_file(make_xlsx(ING, MAT, title_rows=2), "x.xlsx")
    assert not p.errors
    a, w = p.ingredients
    assert a["loss_moisture"] == pytest.approx(.025) and a["price"] == 10500 and a["date"] == "2026-09-30"
    assert w["price"] == 9800 and w["date"] is None and w["origin"] is None   # '9,800원' 파싱
    assert [m["category"] for m in p.materials] == ["pack", "box"]            # '박스' 별칭
    assert p.materials[1]["loss_rate"] == pytest.approx(.03)                  # '3%' 문자열


def test_apply_new_then_unchanged(tmpdb):
    p = importer.parse_file(make_xlsx(ING, MAT), "x.xlsx")
    res = db.bulk_apply(p.ingredients, p.materials, today="2026-10-01")
    assert [r["status"] for r in res] == ["신규"] * 4
    assert db.list_ingredients()[0]["price_per_kg"] in (10500, 9800)
    # 같은 파일을 다시 올리면 전부 변경 없음 (단가 이력도 중복 안 쌓임)
    res2 = db.bulk_apply(p.ingredients, p.materials, today="2026-10-01")
    assert [r["status"] for r in res2] == ["변경없음"] * 4
    ing = next(i for i in db.list_ingredients() if i["name"] == "아몬드")
    assert len(db.price_history(ing["id"])) == 1


def test_update_blank_keeps_existing(tmpdb):
    db.bulk_apply(*[importer.parse_file(make_xlsx(ING, MAT), "x.xlsx").__dict__[k] for k in ("ingredients", "materials")],
                  today="2026-10-01")
    new = [("아몬드", None, None, None, None, None, 11000, "2026-10-15", None)]
    p = importer.parse_file(make_xlsx(new, []), "x.xlsx")
    r = db.bulk_apply(p.ingredients, [], today="2026-10-20")
    assert r[0]["status"] == "수정" and "10,500→11,000" in r[0]["detail"]
    a = next(i for i in db.list_ingredients() if i["name"] == "아몬드")
    assert a["origin"] == "미국" and a["loss_moisture"] == pytest.approx(.025)   # 빈 칸 → 유지
    assert a["price_per_kg"] == 11000 and len(db.price_history(a["id"])) == 2


def test_dry_run_changes_nothing(tmpdb):
    p = importer.parse_file(make_xlsx(ING, MAT), "x.xlsx")
    res = db.bulk_apply(p.ingredients, p.materials, dry_run=True)
    assert all(r["status"] == "신규" for r in res)
    assert db.list_ingredients() == [] and db.list_materials() == []


def test_error_rolls_back_everything(tmpdb):
    mats = [("정상", "포장지", 10, 3, None), ("분류없는신규", None, 10, 3, None)]
    p = importer.parse_file(make_xlsx([], mats), "x.xlsx")
    assert not p.errors                                  # 파싱은 통과, 신규인데 분류 없음은 등록 단계에서 잡음
    res = db.bulk_apply([], p.materials)
    assert any(r["status"] == "오류" for r in res)
    assert db.list_materials() == []                     # '정상' 도 등록되지 않음 (전체 취소)


def test_row_errors_reported_with_location():
    bad_ing = [("아몬드", None, None, 150, None, None, "abc", "2026-13-40", None), (None, "x", None, None, None, None, 1, None, None)]
    bad_mat = [("박스", "이상한분류", 10, 3, None)]
    p = importer.parse_file(make_xlsx(bad_ing, bad_mat), "x.xlsx")
    text = "\n".join(p.errors)
    assert "원료 단가표 2행" in text and "0~100" in text
    assert "원료명이 비어" in text and "3행" in text
    assert "이상한분류" in text


def test_percent_formatted_cells():
    wb = Workbook()
    ws = wb.active
    ws.append(importer.ING_HEADERS)
    ws.append(["아몬드", None, None, 0.05, 0.03, None, 100, None, None])
    ws["D2"].number_format = "0.0%"       # 엑셀에서 % 서식 → 5% 로 해석
    b = io.BytesIO(); wb.save(b)
    p = importer.parse_file(b.getvalue(), "x.xlsx")
    assert p.ingredients[0]["loss_moisture"] == pytest.approx(.05)
    assert p.ingredients[0]["loss_split"] == pytest.approx(.0003)   # 서식 없는 0.03 은 0.03%


def test_fraction_warning():
    rows = [(f"원료{i}", None, None, 0.05, 0.03, 0.01, 100, None, None) for i in range(3)]
    p = importer.parse_file(make_xlsx(rows, []), "x.xlsx")
    assert any("0.05" in w for w in p.warnings)


def test_csv_cp949_and_unknown_file():
    csv_text = "원료명,단가(원/kg),적용일\n아몬드,10500,2026.9.30\n"
    p = importer.parse_file(csv_text.encode("cp949"), "a.csv")
    assert not p.errors and p.ingredients[0]["date"] == "2026-09-30"
    assert importer.parse_file(b"junk", "a.txt").errors
    empty = Workbook(); b = io.BytesIO(); empty.save(b)
    assert importer.parse_file(b.getvalue(), "e.xlsx").errors


def test_export_roundtrip_is_noop(tmpdb):
    p = importer.parse_file(make_xlsx(ING, MAT), "x.xlsx")
    db.bulk_apply(p.ingredients, p.materials, today="2026-10-01")
    blob = importer.build_workbook(db.list_ingredients(), db.list_materials())
    again = importer.parse_file(blob, "export.xlsx")
    assert not again.errors, again.errors
    res = db.bulk_apply(again.ingredients, again.materials, today="2026-10-02")
    assert [r["status"] for r in res] == ["변경없음"] * 4, res


def test_blank_template_has_no_rows():
    wb = load_workbook(io.BytesIO(importer.build_workbook()))
    assert wb.sheetnames == ["작성방법", "원료", "부자재"]
    p = importer.parse_file(importer.build_workbook(), "t.xlsx")
    assert p.errors                              # 빈 양식은 '데이터 없음' (예시 행이 섞여 들어가지 않음)


def test_xls_format():
    xlwt = pytest.importorskip("xlwt")
    wb = xlwt.Workbook()
    ws = wb.add_sheet("단가")
    for c, h in enumerate(importer.ING_HEADERS):
        ws.write(0, c, h)
    pct = xlwt.easyxf(num_format_str="0%")
    ws.write(1, 0, "아몬드"); ws.write(1, 3, 0.05, pct); ws.write(1, 6, 10500)
    b = io.BytesIO(); wb.save(b)
    p = importer.parse_file(b.getvalue(), "old.xls")
    assert not p.errors and p.ingredients[0]["loss_moisture"] == pytest.approx(.05)
    assert p.ingredients[0]["price"] == 10500
