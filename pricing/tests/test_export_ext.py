"""확장 원가 항목 엑셀(1건/목록/전체 데이터) 테스트."""
import io
import pytest
from openpyxl import load_workbook
from pricing import db, export, compare
from pricing.engine import COST_LINES


@pytest.fixture()
def tmpdb(tmp_path, monkeypatch):
    monkeypatch.setattr(db, "DB_PATH", tmp_path / "t.db")
    db.init_db()


def seed():
    db.add_fx_rate("USD", 1300, "2026-01-01")
    db.add_fx_rate("USD", 1400, "2026-06-01")
    a = db.upsert_ingredient("수입A", "미국", "가상", .025, .03, .005, roasting_cost_per_kg=300, sorting_cost_per_kg=200)
    db.add_price(a, None, "2026-01-01", currency="USD", foreign_price=10)
    m1 = db.upsert_material("롤필름", "roll", 25, .03)
    m2 = db.upsert_material("인케이스", "incase", 600, .03)
    pid = db.save_product("제품A", 30, 20, 0, 100, "메모", [{"ingredient_id": a, "ratio": 1.0}],
                          [{"material_id": m1, "qty": 1, "divisor": 1, "per_bag": True},
                           {"material_id": m2, "qty": 1, "divisor": 1}], three_pl_cost=200)
    cid = db.upsert_channel({"name": "가상몰", "kind": "online", "fee_rate": .1, "promo_rate": .02,
                             "parcel_cost": 3000, "event_fixed": 50, "margin_rate": .05})
    return pid, cid


def all_text(wb):
    return " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None)


def test_single_quote_excel_has_new_cost_lines(tmpdb):
    pid, cid = seed()
    q, r = db.calculate_quote(pid, channel_id=cid, as_of="2026-07-01")
    wb = load_workbook(io.BytesIO(export.quotes_to_xlsx([(q, r)], True)))
    txt = all_text(wb)
    for word in ("로스팅비", "3PL 이용료", "롤포장지", "인케이스", "온라인 수수료", "홍보비", "기획전 비용", "택배비",
                 "가상몰", "2026-07-01", "봉수연동", "USD", "1400"):
        assert word in txt, word
    head = [c.value for c in wb["요약"][1]]
    for label, _a, _c in COST_LINES:
        assert label in head
    row = [c.value for c in wb["요약"][2]]
    assert row[head.index("납품가")] == r.price and row[head.index("택배비")] == 3000


def test_list_excel_new_columns(tmpdb):
    pid, cid = seed()
    for ch in (None, cid):
        q, r = db.calculate_quote(pid, channel_id=ch)
        db.save_quote(q, r, "가")
    wb = load_workbook(io.BytesIO(export.quotes_list_to_xlsx(db.search_quotes())))
    head = [c.value for c in wb["견적목록"][1]]
    for col in ("판매처", "견적기준일", "로스팅비", "3PL 이용료", "온라인 수수료", "택배비", "수수료율", "홍보비율"):
        assert col in head
    rows = list(wb["견적목록"].iter_rows(min_row=2, values_only=True))
    assert {r_[head.index("판매처")] for r_ in rows} == {"(기본 설정)", "가상몰"}
    ing_head = [c.value for c in wb["원료상세"][1]]
    assert {"통화", "외화단가", "적용환율(원/1단위)", "로스팅비", "선별비"} <= set(ing_head)
    ing = list(wb["원료상세"].iter_rows(min_row=2, values_only=True))
    assert all(x[ing_head.index("통화")] == "USD" and x[ing_head.index("외화단가")] == 10 for x in ing)


def test_compare_shows_fx_change_as_driver(tmpdb):
    pid, _ = seed()
    qa, ra = db.calculate_quote(pid, as_of="2026-03-01")           # 환율 1300
    qb, rb = db.calculate_quote(pid, as_of="2026-07-01")           # 환율 1400, 현지가 동일
    meta = lambda i: {"id": i, "created_at": "", "created_by": "", "product_name": "제품A", "price": 0, "memo": "",
                      "mode": "자동 계산"}
    c = compare.compare_quotes((qa, ra, meta(1)), (qb, rb, meta(2)))
    assert any("환율 1,300.00→1,400.00" in d for d in c["drivers"]), c["drivers"]
    rules = {r["항목"]: r for r in c["rules"]}
    assert rules["견적 기준일"]["다름"] == "●"
    ing = c["ingredients"][0]
    assert ing["환율 A"] == 1300 and ing["환율 B"] == 1400 and ing["외화단가 A"] == ing["외화단가 B"] == 10
    load_workbook(io.BytesIO(export.compare_to_xlsx(c)))


def test_all_data_excel_has_every_sheet(tmpdb):
    pid, cid = seed()
    q, r = db.calculate_quote(pid, channel_id=cid)
    db.save_quote(q, r, "가", "메모")
    wb = load_workbook(io.BytesIO(export.all_data_xlsx(db.collect_all_data())))
    assert wb.sheetnames == ["원료 마스터", "원료 단가이력", "부자재 마스터", "부자재 단가이력", "환율 이력", "통화",
                             "판매처", "제품", "제품 구성(BOM)", "견적 이력 요약"]
    assert wb["원료 단가이력"].max_row == 2 and wb["환율 이력"].max_row == 3
    assert wb["부자재 단가이력"].max_row == 3 and wb["제품 구성(BOM)"].max_row == 4
    txt = all_text(wb)
    assert "가상몰" in txt and "봉수 연동" in txt and "(기타)" not in txt          # 기본 통화만 있으니 (기타) 없음
    db.add_currency("EUR", "유로", "유럽연합")
    assert "(기타)" in all_text(load_workbook(io.BytesIO(export.all_data_xlsx(db.collect_all_data()))))
