import copy
import io
import pytest
from openpyxl import load_workbook
from pricing import db
from pricing.compare import compare_quotes
from pricing.engine import (QuoteInput, IngredientLine as I, MaterialLine as M, calculate,
                            calc_mode, MODE_AUTO, MODE_FIXED)
from pricing.export import quotes_list_to_xlsx, quotes_to_xlsx


def make(price_a=10000, ratio_a=.6, **kw):
    d = dict(product_name="테스트", bag_count=10, unit_weight_g=20,
             ingredients=[I("A", ratio_a, price_a, .05, "미국", "가상상사", "2026-09-01", "마스터 단가"),
                          I("B", 1 - ratio_a, 20000, .02)],
             materials=[M("소포장", "pack", 10, 10, 1, .03), M("박스", "box", 1000, 1, 10, .03)],
             shipping_cost=100, sga_rate=.05, rounding="round")
    d.update(kw)
    q = QuoteInput(**d)
    return q, calculate(q), {"id": 1, "created_at": "2026-10-01 10:00", "created_by": "홍길동",
                             "product_name": q.product_name, "price": 0, "memo": "", "mode": calc_mode(q)}


# ---------- 비교 ----------
def test_compare_price_driver_and_summary():
    a, b = make(price_a=10000), make(price_a=11000)
    c = compare_quotes(a, b)
    assert c["same_product"]
    assert c["summary"][-2]["항목"] == "실질 마진"
    price_row = next(r for r in c["summary"] if r["항목"] == "납품가")
    assert price_row["차이"] == b[1].price - a[1].price > 0
    top = c["drivers"][0]
    assert "원료 A" in top and "10,000→11,000" in top
    assert "→" in c["headline"] and "%" in c["headline"]
    assert next(r for r in c["ingredients"] if r["원료"] == "A")["상태"] == "변경"
    assert next(r for r in c["ingredients"] if r["원료"] == "B")["상태"] == "동일"


def test_compare_added_removed_and_rules():
    a = make()
    q2, r2, m2 = make(sga_rate=.0, fixed_price=1000)
    q2.ingredients = q2.ingredients[:1] + [I("C", .4, 5000, 0)]
    q2.materials = q2.materials[:1]
    r2 = calculate(q2)
    c = compare_quotes(a, (q2, r2, m2))
    st = {r["원료"]: r["상태"] for r in c["ingredients"]}
    assert st["B"] == "삭제" and st["C"] == "추가"
    assert {r["부자재"]: r["상태"] for r in c["materials"]}["박스"] == "삭제"
    rules = {r["항목"]: r for r in c["rules"]}
    assert rules["계산 방식"]["다름"] == "●" and rules["판관비율"]["다름"] == "●"
    assert rules["봉수"]["다름"] == ""


# ---------- 계산 방식 / 출처 기록 ----------
def test_mode_and_sources_roundtrip():
    q, r, _ = make()
    assert calc_mode(q) == MODE_AUTO
    assert calc_mode(make(fixed_price=900)[0]) == MODE_FIXED
    row = r.ingredient_rows[0]
    assert (row["origin"], row["supplier"], row["price_date"]) == ("미국", "가상상사", "2026-09-01")


@pytest.fixture()
def tmpdb(tmp_path, monkeypatch):
    monkeypatch.setattr(db, "DB_PATH", tmp_path / "t.db")
    db.init_db()


def seed():
    a = db.upsert_ingredient("아몬드", "미국", "가상상사", .025, .03, .0054)
    db.add_price(a, 10000, "2026-01-01")
    db.add_price(a, 11000, "2026-06-01")
    h = db.upsert_ingredient("호두", "칠레", "다른상사", 0, .03, .005)
    db.add_price(h, 9000, "2026-02-01")
    m = db.upsert_material("카톤", "box", 1000, .03)
    return db.save_product("제품1", 10, 20, 300, 100, "",
                           [{"ingredient_id": a, "ratio": .5},
                            {"ingredient_id": h, "ratio": .5, "price_override": 9500}],
                           [{"material_id": m, "qty": 1, "divisor": 10}])


def test_quote_records_source_info(tmpdb):
    pid = seed()
    q, r = db.calculate_quote(pid)
    by = {i.name: i for i in q.ingredients}
    assert (by["아몬드"].origin, by["아몬드"].supplier, by["아몬드"].price_date, by["아몬드"].price_source) == \
        ("미국", "가상상사", "2026-06-01", "마스터 단가")
    assert by["호두"].price_source == "제품 직접입력" and by["호두"].price_date == ""
    qid = db.save_quote(q, r, "tester")
    q2, r2, meta = db.load_quote(qid)
    assert q2.ingredients[0].price_date == "2026-06-01" and meta["mode"] == MODE_AUTO


def test_old_quote_without_new_fields_still_loads(tmpdb):
    """새 필드가 없던 예전 이력도 열려야 한다."""
    import json
    from dataclasses import asdict
    pid = seed()
    q, r = db.calculate_quote(pid)
    d = asdict(q)
    for i in d["ingredients"]:
        for k in ("origin", "supplier", "price_date", "price_source"):
            i.pop(k)
    rd = r.to_dict()
    for x in rd["ingredient_rows"]:
        for k in ("origin", "supplier", "price_date", "price_source"):
            x.pop(k)
    with db.connect() as con:
        con.execute("INSERT INTO quotes(product_name,created_at,created_by,memo,price,input_json,result_json) "
                    "VALUES('구버전','2026-01-01 00:00','x','',1,?,?)",
                    (json.dumps(d, ensure_ascii=False), json.dumps(rd, ensure_ascii=False)))
    q2, r2, meta = db.search_quotes()[0]
    assert q2.ingredients[0].price_date == "" and meta["mode"] == MODE_AUTO
    assert quotes_list_to_xlsx([(q2, r2, meta)])          # 엑셀도 오류 없이 생성


# ---------- 검색 / 목록 엑셀 ----------
def test_search_filters(tmpdb):
    pid = seed()
    for who, fixed in (("가", None), ("나", 5000.0), ("가", None)):
        q, r = db.calculate_quote(pid, fixed_price=fixed)
        db.save_quote(q, r, who)
    assert len(db.search_quotes()) == 3
    assert len(db.search_quotes(users=["가"])) == 2
    assert len(db.search_quotes(modes=[MODE_FIXED])) == 1
    assert len(db.search_quotes(products=["없는제품"])) == 0
    today = db.search_quotes()[0][2]["created_at"][:10]
    assert len(db.search_quotes(date_from=today, date_to=today)) == 3
    assert len(db.search_quotes(date_from="2999-01-01")) == 0
    opts = db.quote_filter_options()
    assert opts["products"] == ["제품1"] and opts["users"] == ["가", "나"]


def test_list_excel_contents(tmpdb):
    pid = seed()
    q, r = db.calculate_quote(pid)
    db.save_quote(q, r, "가", "메모1")
    q, r = db.calculate_quote(pid, fixed_price=5000)
    db.save_quote(q, r, "나", "지정가")
    items = db.search_quotes()
    wb = load_workbook(io.BytesIO(quotes_list_to_xlsx(items)))
    assert wb.sheetnames == ["견적목록", "원료상세", "부자재상세"]
    head = [c.value for c in wb["견적목록"][1]]
    for col in ("견적번호", "일시", "작성자", "제품", "계산방식", "납품가", "실질마진율", "경고", "메모"):
        assert col in head
    rows = list(wb["견적목록"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2
    assert {row[head.index("계산방식")] for row in rows} == {MODE_AUTO, MODE_FIXED}
    ing_head = [c.value for c in wb["원료상세"][1]]
    ing = list(wb["원료상세"].iter_rows(min_row=2, values_only=True))
    assert len(ing) == 4 and "단가기준일" in ing_head and "공급처" in ing_head
    assert any(r[ing_head.index("단가출처")] == "제품 직접입력" for r in ing)


def test_single_quote_excel_has_meta_and_mode(tmpdb):
    pid = seed()
    q, r = db.calculate_quote(pid, fixed_price=5000)
    qid = db.save_quote(q, r, "가", "m")
    q2, r2, meta = db.load_quote(qid)
    wb = load_workbook(io.BytesIO(quotes_to_xlsx([(q2, r2, meta)], True)))
    text = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value)
    assert MODE_FIXED in text and "가상상사" in text and "2026-06-01" in text


# ---------- 삭제 잠금 ----------
def test_delete_locked_by_default_and_audited(tmpdb):
    pid = seed()
    q, r = db.calculate_quote(pid)
    qid = db.save_quote(q, r, "가")
    with pytest.raises(PermissionError):
        db.delete_quote(qid, user="나")
    assert len(db.search_quotes()) == 1                    # 삭제 안 됨
    s = db.get_settings(); s["allow_quote_delete"] = True
    db.save_settings(s)
    db.delete_quote(qid, user="나")
    assert db.search_quotes() == []
    log = db.list_audit()
    assert log[0]["action"] == "견적 삭제" and log[0]["user"] == "나" and "제품1" in log[0]["detail"]
