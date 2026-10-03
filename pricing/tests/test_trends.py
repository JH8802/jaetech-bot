"""가격 추이·환율 분해·연도별·제품 원가 추이·엑셀 (가짜 데이터)."""
import io
import pytest
from openpyxl import load_workbook
from pricing import db, trends, export
from pricing.prices import PriceBook


def book():
    cur = [{"code": "USD", "quote_unit": 1}, {"code": "VND", "quote_unit": 100}]
    fx = [{"id": 1, "currency": "USD", "rate": 1300.0, "rate_date": "2025-01-01"},
          {"id": 2, "currency": "USD", "rate": 1400.0, "rate_date": "2026-01-10"},
          {"id": 3, "currency": "USD", "rate": 1500.0, "rate_date": "2026-02-01"}]
    ing = [  # 1: 달러 단가(자동환율) 10 → 11(2026-01-20)   2: 원화 단가 9000 → 9900(2026-01-05)
        {"id": 1, "ingredient_id": 1, "price_per_kg": 13000.0, "effective_date": "2025-01-01",
         "currency": "USD", "foreign_price": 10.0, "fx_mode": "auto"},
        {"id": 2, "ingredient_id": 1, "price_per_kg": 15400.0, "effective_date": "2026-01-20",
         "currency": "USD", "foreign_price": 11.0, "fx_mode": "auto"},
        {"id": 3, "ingredient_id": 2, "price_per_kg": 9000.0, "effective_date": "2025-01-01", "currency": "KRW"},
        {"id": 4, "ingredient_id": 2, "price_per_kg": 9900.0, "effective_date": "2026-01-05", "currency": "KRW"}]
    return PriceBook(cur, fx, ing, [])


def test_daily_series_daily_changes_and_flags():
    s = trends.daily_series(book(), "ingredient", 1, "2026-01-08", "2026-02-02")
    by = {r["date"]: r for r in s}
    assert len(s) == 26
    assert by["2026-01-09"]["krw"] == pytest.approx(13000)           # 1300 × 10
    assert by["2026-01-10"]["krw"] == pytest.approx(14000) and by["2026-01-10"]["changed"]   # 환율만 변동
    assert by["2026-01-19"]["krw"] == pytest.approx(14000)
    assert by["2026-01-20"]["krw"] == pytest.approx(15400) and by["2026-01-20"]["changed"]   # 현지 단가 변동
    assert by["2026-02-01"]["krw"] == pytest.approx(16500) and by["2026-02-01"]["changed"]   # 환율 1500 × 11


def test_series_skips_dates_before_first_record():
    assert trends.daily_series(book(), "ingredient", 1, "2024-12-30", "2025-01-02")[0]["date"] == "2025-01-01"
    assert trends.daily_series(book(), "ingredient", 99, "2026-01-01", "2026-01-05") == []


def test_period_stats():
    s = trends.daily_series(book(), "ingredient", 2, "2026-01-01", "2026-01-10")
    st = trends.period_stats(s)
    assert st["start"] == 9000 and st["end"] == 9900 and st["rate"] == pytest.approx(0.10)
    assert st["min"] == 9000 and st["max"] == 9900 and st["max_date"] == "2026-01-05"
    assert st["change_days"] == 1 and st["days"] == 10
    assert trends.period_stats([]) is None


def test_decompose_effects_sum_to_total():
    s = trends.daily_series(book(), "ingredient", 1, "2026-01-01", "2026-02-05")
    d = trends.decompose(s)
    assert d["currency"] == "USD"
    assert d["foreign_rate"] == pytest.approx(0.10) and d["fx_rate"] == pytest.approx(1500 / 1300 - 1)
    assert d["price_effect"] == pytest.approx((11 - 10) * 1300)        # 현지가 효과 = 단가변동 × 기초환율
    assert d["fx_effect"] == pytest.approx(11 * (1500 - 1300))          # 환율 효과 = 기말 현지가 × 환율변동
    assert d["price_effect"] + d["fx_effect"] == pytest.approx(d["total"])
    assert d["index"][0]["원화 단가"] == pytest.approx(100) and d["index"][-1]["환율"] == pytest.approx(1500 / 13)
    assert trends.decompose(trends.daily_series(book(), "ingredient", 2, "2026-01-01", "2026-01-10")) is None


def test_aggregate_weekly_monthly():
    s = trends.daily_series(book(), "ingredient", 2, "2026-01-01", "2026-02-28")
    m = trends.aggregate(s, "M")
    assert [x["period"] for x in m] == ["2026-01", "2026-02"]
    assert m[1]["avg"] == 9900 and m[0]["min"] == 9000 and m[0]["max"] == 9900
    w = trends.aggregate(s, "W")
    assert all(len(x["period"]) == 10 for x in w) and len(trends.aggregate(s, "D")) == len(s)


def test_ranking_sorted_by_rate_and_fx_columns():
    rows = trends.ranking_table(book(), "ingredient", {1: "수입", 2: "국내"}, "2026-01-01", "2026-02-05")
    assert [r["name"] for r in rows] == ["수입", "국내"] or [r["name"] for r in rows] == ["국내", "수입"]
    assert rows[0]["rate"] >= rows[1]["rate"]
    imp = next(r for r in rows if r["name"] == "수입")
    dom = next(r for r in rows if r["name"] == "국내")
    assert imp["fx_rate"] is not None and dom["fx_rate"] is None


def test_yearly_table_yoy():
    rows = trends.yearly_table(book(), "ingredient", {2: "국내"}, [2025, 2026], today="2026-06-30")
    y25 = next(r for r in rows if r["year"] == 2025)
    y26 = next(r for r in rows if r["year"] == 2026)
    assert y25["base_label"] == "최초 등록가" and y25["rate"] == pytest.approx(0)
    assert y26["base"] == 9000 and y26["end"] == 9900 and y26["rate"] == pytest.approx(0.10)
    assert y26["partial"] and not y25["partial"] and y26["base_label"] == "2025년 말"


def test_yearly_table_fx_drives_krw_change():
    rows = trends.yearly_table(book(), "ingredient", {1: "수입"}, [2026], today="2026-12-31")
    r = rows[0]
    assert r["base"] == pytest.approx(13000) and r["end"] == pytest.approx(11 * 1500)
    assert r["rate"] == pytest.approx(16500 / 13000 - 1)


def test_fx_daily_uses_quote_unit():
    pb = PriceBook([{"code": "VND", "quote_unit": 100}], [{"id": 1, "currency": "VND", "rate": 5.5,
                   "rate_date": "2026-01-02"}], [], [])
    s = trends.fx_daily(pb, "VND", "2026-01-01", "2026-01-04")
    assert [r["date"] for r in s] == ["2026-01-02", "2026-01-03", "2026-01-04"]
    assert s[0]["rate"] == pytest.approx(5.5)                 # 100동당 5.5원으로 표기


def test_range_validation():
    with pytest.raises(ValueError):
        list(trends.daterange("2026-02-01", "2026-01-01"))
    with pytest.raises(ValueError):
        list(trends.daterange("2000-01-01", "2026-01-01"))
    d = trends.sample_dates("2026-01-15", "2026-04-10", "M")
    assert d == ["2026-01-15", "2026-01-31", "2026-02-28", "2026-03-31", "2026-04-10"]
    assert trends.sample_dates("2026-01-01", "2026-01-15", "W") == ["2026-01-01", "2026-01-08", "2026-01-15"]


# ---------------- DB 연동: 제품 원가 추이 ----------------
@pytest.fixture()
def tmpdb(tmp_path, monkeypatch):
    monkeypatch.setattr(db, "DB_PATH", tmp_path / "t.db")
    db.init_db()


def test_product_series_follows_fx_and_price(tmpdb):
    db.add_fx_rate("USD", 1300, "2026-01-01")
    db.add_fx_rate("USD", 1500, "2026-03-01")
    a = db.upsert_ingredient("수입A")
    db.add_price(a, None, "2026-01-01", currency="USD", foreign_price=10)
    b = db.upsert_ingredient("국내B")
    db.add_price(b, 8000, "2026-01-01")
    db.add_price(b, 9000, "2026-04-01")
    pid = db.save_product("p", 10, 20, 0, 0, "", [{"ingredient_id": a, "ratio": .5}, {"ingredient_id": b, "ratio": .5}], [])
    pb = db.load_pricebook()
    rows = trends.product_series(db.assemble_quote, db.get_product_detail(pid), db.get_settings(), None, pb,
                                 "2026-01-31", "2026-04-30", "M")
    by = {r["date"][:7]: r for r in rows}
    assert by["2026-01"]["price"] < by["2026-03"]["price"] < by["2026-04"]["price"]     # 환율↑, 그다음 국내단가↑
    # 3월(환율 1500, 국내 8000)과 2월(환율 1300) 차이 = 달러 원료 원가 차이만큼
    kg = 20 * .5 * 10 / 1000
    assert by["2026-03"]["direct_cost"] - by["2026-02"]["direct_cost"] == pytest.approx(10 * (1500 - 1300) * kg)


# ---------------- 엑셀 ----------------
def test_excel_exports_have_data_and_charts():
    s = trends.daily_series(book(), "ingredient", 1, "2026-01-01", "2026-02-05")
    st, dec = trends.period_stats(s), trends.decompose(s)
    blob = export.item_trend_xlsx("수입", "원료", s, st, dec, trends.aggregate(s, "M"), "월별")
    wb = load_workbook(io.BytesIO(blob))
    assert wb.sheetnames[0] == "요약" and "일별 가격" in wb.sheetnames and "지수(기준=100)" in wb.sheetnames
    ws = wb["일별 가격"]
    assert ws.max_row == len(s) + 1 and len(ws._charts) == 1
    text = {str(c.value) for row in wb["요약"].iter_rows() for c in row}
    assert "환율 변동 효과(원)" in text

    rk = trends.ranking_table(book(), "ingredient", {1: "수입", 2: "국내"}, "2026-01-01", "2026-02-05")
    wb = load_workbook(io.BytesIO(export.ranking_xlsx(rk, "원료", "2026-01-01", "2026-02-05")))
    assert wb.active.max_row == 3 and len(wb.active._charts) == 1

    yr = trends.yearly_table(book(), "ingredient", {1: "수입", 2: "국내"}, [2025, 2026], today="2026-12-31")
    wb = load_workbook(io.BytesIO(export.yearly_xlsx(yr, "원료")))
    assert wb.sheetnames == ["원료 연도별 변동", "변동률 표"]

    fxs = {"USD": trends.fx_daily(book(), "USD", "2026-01-01", "2026-01-31")}
    wb = load_workbook(io.BytesIO(export.fx_xlsx(fxs, {"USD": "미국 달러"})))
    assert wb["USD 환율"].max_row == 32 and len(wb["USD 환율"]._charts) == 1

    empty = load_workbook(io.BytesIO(export.ranking_xlsx([], "원료", "a", "b")))
    assert empty.active["A2"].value == "(조회 결과 없음)"
