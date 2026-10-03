"""화면 스모크 테스트: 모든 메뉴가 예외 없이 열리고 핵심 흐름이 동작하는지 (가짜 데이터)."""
from pathlib import Path

import pytest

pytest.importorskip("streamlit")
from streamlit.testing.v1 import AppTest  # noqa: E402

from pricing import db  # noqa: E402

APP = str(Path(__file__).resolve().parents[1] / "app.py")
PAGES = ["견적 계산", "제품 관리", "원료 단가", "부자재", "환율", "판매처", "가격 추이",
         "엑셀 일괄 등록", "견적 이력", "설정"]


@pytest.fixture()
def app_db(tmp_path, monkeypatch):
    monkeypatch.setattr(db, "DB_PATH", tmp_path / "ui.db")
    db.init_db()
    for d, r in [("2026-01-02", 1300), ("2026-03-01", 1350), ("2026-06-01", 1420)]:
        db.add_fx_rate("USD", r, d)
    db.add_fx_rate("VND", 5.5, "2026-01-02")
    a = db.upsert_ingredient("수입A(가상)", "미국", "가상", .025, .03, .005, 300, 200)
    db.add_price(a, None, "2026-01-02", currency="USD", foreign_price=7.0)
    db.add_price(a, None, "2026-05-01", currency="USD", foreign_price=7.5)
    h = db.upsert_ingredient("국내B(가상)")
    db.add_price(h, 9000, "2026-01-02")
    db.add_price(h, 9900, "2026-07-01")
    m1 = db.upsert_material("롤필름", "roll", None, .03)
    db.add_material_price(m1, None, "2026-01-02", currency="VND", foreign_price=500)
    m2 = db.upsert_material("카톤", "box", 1000, .03)
    pid = db.save_product("제품(가상)", 30, 20, 0, 100, "",
                          [{"ingredient_id": a, "ratio": .5}, {"ingredient_id": h, "ratio": .5}],
                          [{"material_id": m1, "qty": 1, "divisor": 1, "per_bag": True},
                           {"material_id": m2, "qty": 1, "divisor": 10}], three_pl_cost=200)
    db.upsert_channel({"name": "가상마트", "kind": "offline", "logistics_rate": .03, "logistics_fixed": 100})
    db.upsert_channel({"name": "가상몰", "kind": "online", "fee_rate": .12, "promo_rate": .02, "event_rate": .01,
                       "parcel_cost": 3000, "margin_rate": .05})
    return pid


def open_page(page, who="홍길동"):
    at = AppTest.from_file(APP, default_timeout=90).run()
    at.sidebar.text_input[0].set_value(who).run()
    at.sidebar.radio[0].set_value(page).run()
    return at


@pytest.mark.parametrize("page", PAGES)
def test_every_page_opens_without_exception(app_db, page):
    at = open_page(page)
    assert not at.exception, [str(e.value)[:200] for e in at.exception]


def test_sidebar_menu_order_and_new_menus(app_db):
    at = AppTest.from_file(APP, default_timeout=90).run()
    assert list(at.sidebar.radio[0].options) == PAGES


def test_quote_with_channel_and_save(app_db):
    at = open_page("견적 계산")
    base = at.metric[0].value
    online = [o for o in at.selectbox[1].options if "가상몰" in o][0]
    at.selectbox[1].set_value(online).run()
    assert at.metric[0].value != base and not at.exception
    [b for b in at.button if "이력에 저장" in b.label][0].click().run()
    assert any("저장됨" in s.value for s in at.success)
    q, r, meta = db.search_quotes()[0]
    assert q.channel_name == "가상몰" and q.as_of_date and r.parcel == 3000 and meta["created_by"] == "홍길동"


def test_trends_page_draws_charts_and_downloads(app_db):
    at = open_page("가격 추이")
    kinds = {getattr(e, "type", "") for e in at.main}
    assert "vega_lite_chart" in kinds and "download_button" in kinds


def test_fx_page_saves_rate_for_selected_currency_and_adds_other_currency(app_db):
    at = open_page("환율")
    sb = [x for x in at.selectbox if x.label == "통화"][0]
    assert sb.options[0] == "미국 ($)" and sb.options[1] == "베트남 (₫)" and sb.options[2] == "튀르키예 (₺)"
    sb.set_value(sb.options[0])
    [n for n in at.number_input if n.label.startswith("환율")][0].set_value(1450.0)
    at.run()
    [b for b in at.button if "환율 저장" in b.label][0].click().run()
    assert db.list_fx_rates("USD")[0]["rate"] == 1450.0 and not at.exception
    for ti in at.text_input:
        if ti.label.startswith("통화 코드"):
            ti.set_value("eur")
        elif ti.label == "통화 이름":
            ti.set_value("유로")
        elif ti.label == "국가":
            ti.set_value("유럽연합")
        elif ti.label == "통화 기호":
            ti.set_value("€")
    at.run()
    [b for b in at.button if b.label == "통화 추가"][0].click().run()
    eur = next(c for c in db.list_currencies() if c["code"] == "EUR")
    assert eur["symbol"] == "€" and db.currency_label(eur) == "(기타) 유럽연합 (€)"


def test_fx_history_help_explains_data_source(app_db):
    at = open_page("환율")
    helps = " ".join(getattr(e, "help", "") or "" for e in at.subheader)
    assert "인터넷에서 자동으로 가져오지 않습니다" in helps and "직전에 입력한 환율" in helps


def _fill(at, label_startswith, value):
    for n in at.number_input:
        if n.label.startswith(label_startswith):
            n.set_value(value)
            return
    raise AssertionError(f"number_input not found: {label_startswith}")


def _text(at, label, value):
    [t for t in at.text_input if t.label == label][0].set_value(value)


def _click(at, label_part):
    [b for b in at.button if label_part in b.label][0].click().run()


def test_ingredient_form_has_base_price_before_roasting_and_fx_options(app_db):
    at = open_page("원료 단가")
    labels = [n.label for n in at.number_input]
    assert "기본 원료 단가(원/kg)" in labels
    assert labels.index("기본 원료 단가(원/kg)") < labels.index("로스팅비(원/kg)")        # 로스팅비 앞쪽
    fx = [x for x in at.selectbox if x.label.startswith("환율 (수입")][0]
    assert list(fx.options) == ["해당 없음 (원화 ₩)", "미국 ($)", "베트남 (₫)", "튀르키예 (₺)", "(기타) 새 통화 직접 입력"]


def test_new_domestic_ingredient_with_base_price(app_db):
    at = open_page("원료 단가")
    _text(at, "원료명", "국산땅콩(가상)")
    _fill(at, "기본 원료 단가", 5200.0)
    at.run()
    _click(at, "원료 저장")
    assert not at.exception and not at.error
    ing = next(i for i in db.list_ingredients() if i["name"] == "국산땅콩(가상)")
    assert ing["currency"] == "KRW" and ing["price_per_kg"] == 5200.0


def test_new_foreign_ingredient_with_typed_fx_rate_is_fixed(app_db):
    at = open_page("원료 단가")
    _text(at, "원료명", "수입피스타치오")
    fx = [x for x in at.selectbox if x.label.startswith("환율 (수입")][0]
    fx.set_value("미국 ($)").run()
    rate_input = [n for n in at.number_input if n.label.startswith("환율 (원 /")][0]
    assert rate_input.value == 1420.0                      # 환율 메뉴의 최근 값이 기본으로 채워짐 (고쳐 쓸 수 있음)
    assert [n for n in at.number_input if n.label.startswith("기본 원료 단가")] == []   # 외화 선택 시 원화칸은 자동계산(잠금)
    _fill(at, "외화 단가", 8.0)
    _fill(at, "환율 (원 /", 1400.0)                         # 단가를 받은 당시 환율을 직접 입력
    at.run()
    assert any("11,200.00원" in c.value for c in at.caption)  # 8$ × 1,400원 = 11,200원 미리보기
    _click(at, "원료 저장")
    assert not at.exception and not at.error
    ing = next(i for i in db.list_ingredients() if i["name"] == "수입피스타치오")
    assert (ing["currency"], ing["foreign_price"], ing["fx_mode"], ing["fx_fixed_rate"]) == ("USD", 8.0, "fixed", 1400.0)
    assert ing["price_per_kg"] == pytest.approx(11200.0)   # 환율 메뉴(1420)가 아니라 입력한 1400으로 환산


def test_vnd_rate_is_per_100_dong(app_db):
    at = open_page("원료 단가")
    _text(at, "원료명", "베트남캐슈")
    [x for x in at.selectbox if x.label.startswith("환율 (수입")][0].set_value("베트남 (₫)").run()
    labels = [n.label for n in at.number_input]
    assert any("외화 단가 (₫/kg)" in l for l in labels) and any("환율 (원 / 100₫)" in l for l in labels)
    _fill(at, "외화 단가", 250000.0)
    _fill(at, "환율 (원 /", 5.6)
    at.run()
    _click(at, "원료 저장")
    ing = next(i for i in db.list_ingredients() if i["name"] == "베트남캐슈")
    assert ing["price_per_kg"] == pytest.approx(250000 * 5.6 / 100)


def test_other_currency_new_entry_creates_currency_with_symbol(app_db):
    at = open_page("원료 단가")
    _text(at, "원료명", "유럽올리브")
    [x for x in at.selectbox if x.label.startswith("환율 (수입")][0].set_value("(기타) 새 통화 직접 입력").run()
    _text(at, "통화 코드 (영문 2~6자)", "eur")
    _text(at, "통화 이름", "유로")
    _text(at, "국가", "유럽연합")
    _text(at, "통화 기호", "€")
    _fill(at, "외화 단가", 6.0)
    _fill(at, "환율 (원 /", 1500.0)
    at.run()
    _click(at, "원료 저장")
    assert not at.exception and not at.error, [e.value for e in at.error]
    eur = next(c for c in db.list_currencies() if c["code"] == "EUR")
    assert eur["symbol"] == "€"
    ing = next(i for i in db.list_ingredients() if i["name"] == "유럽올리브")
    assert ing["currency"] == "EUR" and ing["price_per_kg"] == pytest.approx(9000.0)


def test_editing_master_only_does_not_add_price_history(app_db):
    at = open_page("원료 단가")
    at.selectbox(key="ing_pick").set_value("국내B(가상)").run()
    before = len(db.price_history(next(i for i in db.list_ingredients() if i["name"] == "국내B(가상)")["id"]))
    [n for n in at.number_input if n.label == "수분 loss(%)"][0].set_value(4.0)
    at.run()
    _click(at, "원료 저장")
    ing = next(i for i in db.list_ingredients() if i["name"] == "국내B(가상)")
    assert ing["loss_moisture"] == pytest.approx(0.04)
    assert len(db.price_history(ing["id"])) == before        # 단가는 그대로 → 이력 줄이 늘지 않음


def test_changing_price_adds_dated_history_row(app_db):
    at = open_page("원료 단가")
    at.selectbox(key="ing_pick").set_value("국내B(가상)").run()
    _fill(at, "기본 원료 단가", 10500.0)
    at.run()
    _click(at, "원료 저장")
    ing = next(i for i in db.list_ingredients() if i["name"] == "국내B(가상)")
    assert ing["price_per_kg"] == 10500.0 and len(db.price_history(ing["id"])) == 3


def test_material_form_has_same_fx_options_and_saves_foreign_price(app_db):
    at = open_page("부자재")
    fx = [x for x in at.selectbox if x.label.startswith("환율 (수입")][0]
    assert list(fx.options)[:4] == ["해당 없음 (원화 ₩)", "미국 ($)", "베트남 (₫)", "튀르키예 (₺)"]
    assert "기본 부자재 단가(원)" in [n.label for n in at.number_input]
    _text(at, "부자재명", "수입스티커")
    fx.set_value("튀르키예 (₺)").run()
    _fill(at, "외화 단가", 2.0)
    _fill(at, "환율 (원 /", 33.0)
    at.run()
    _click(at, "부자재 저장")
    assert not at.exception and not at.error
    m = next(x for x in db.list_materials() if x["name"] == "수입스티커")
    assert (m["currency"], m["foreign_price"], m["fx_mode"]) == ("TRY", 2.0, "fixed") and m["unit_price"] == pytest.approx(66.0)


def test_help_icons_present_on_key_pages(app_db):
    def helps(page):
        at = open_page(page)
        els = list(at.header) + list(at.subheader) + list(at.number_input) + list(at.selectbox) + list(at.metric)
        return [e for e in els if getattr(e, "help", None)]
    for page in PAGES:
        assert len(helps(page)) >= 1, page                       # 메뉴마다 (?) 설명이 하나 이상
    assert len(helps("원료 단가")) >= 8 and len(helps("견적 계산")) >= 6 and len(helps("가격 추이")) >= 6


def test_bulk_page_with_fx_sheet(app_db, monkeypatch):
    import io
    import streamlit as st
    from openpyxl import Workbook
    from pricing import importer
    wb = Workbook()
    wi = wb.active; wi.title = "원료"; wi.append(importer.ING_HEADERS)
    wi.append(["신규수입원료", None, None, None, None, None, None, "2026-02-01", None, None, None, "USD", 5.0, None])
    wb.create_sheet("부자재").append(importer.MAT_HEADERS)
    wf = wb.create_sheet("환율"); wf.append(importer.FX_HEADERS); wf.append(["USD", "2026-02-01", 1310])
    buf = io.BytesIO(); wb.save(buf)

    class Fake:
        name = "x.xlsx"
        def getvalue(self): return buf.getvalue()
    monkeypatch.setattr(st, "file_uploader", lambda *a, **k: Fake())
    at = open_page("엑셀 일괄 등록")
    assert not at.exception and [m.label for m in at.metric] == ["신규", "수정", "변경없음", "오류"]
    [b for b in at.button if "등록 실행" in b.label][0].click().run()
    assert any("등록 완료" in s.value for s in at.success)
    assert any(i["name"] == "신규수입원료" for i in db.list_ingredients())
    assert any(r["rate_date"] == "2026-02-01" and r["rate"] == 1310 for r in db.list_fx_rates("USD"))


def _charts(at):
    return [e for e in at.main if getattr(e, "type", "") == "vega_lite_chart"]


def test_trends_bar_views_and_fx_decomposition(app_db):
    ing = {i["name"]: i["id"] for i in db.list_ingredients()}
    at = open_page("가격 추이")
    at.selectbox(key="tr_item1").set_value(ing["수입A(가상)"]).run()          # 달러 수입 원료
    assert not at.exception
    assert any("변동 원인 분해" in m.value for m in at.markdown)               # 현지가 × 환율 분해
    assert any("환율 변동 효과" in c.value for c in at.caption)
    n_line = len(_charts(at))
    assert n_line >= 2                                                       # 일별 선 + 지수 선
    for label in ("막대 그래프 (주별 평균)", "막대 그래프 (월별 평균)"):
        at.radio(key="tr_view1").set_value(label).run()
        assert not at.exception and len(_charts(at)) >= 2


def test_trends_ranking_yearly_and_product_tabs_render(app_db):
    at = open_page("가격 추이")
    assert not at.exception
    # 탭 4개 모두 렌더링됨: 변동률 순위(상승/하락 지표), 연도별 표, 제품 원가 추이 지표
    labels = [m.label for m in at.metric]
    assert "상승 품목" in labels and "시작 납품가" in labels
    assert len(at.tabs) == 4
