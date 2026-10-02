import pytest
from pricing import db
from pricing.export import quotes_to_xlsx
from openpyxl import load_workbook
from io import BytesIO


@pytest.fixture()
def tmpdb(tmp_path, monkeypatch):
    monkeypatch.setattr(db, "DB_PATH", tmp_path / "t.db")
    db.init_db()
    return tmp_path


def seed():
    a = db.upsert_ingredient("아몬드", "미국", "가상공급처", .025, .03, .0054)
    db.add_price(a, 10000, "2026-01-01")
    db.add_price(a, 11000, "2026-06-01", "인상")
    m1 = db.upsert_material("소포장지", "pack", 20, .03)
    m2 = db.upsert_material("카톤", "box", 1000, .03)
    return db.save_product("테스트제품", 10, 20,
                           300, 100, "", [{"ingredient_id": a, "ratio": 1.0}],
                           [{"material_id": m1, "qty": 10, "divisor": 1},
                            {"material_id": m2, "qty": 1, "divisor": 10}]), a


def test_latest_price_and_master_loss(tmpdb):
    pid, a = seed()
    q = db.build_quote_input(pid)
    assert q.ingredients[0].price_per_kg == 11000            # 최신 단가
    assert q.ingredients[0].loss_rate == pytest.approx(.0604)  # 수분+소분+선별


def test_price_change_reflects_in_quote(tmpdb):
    pid, a = seed()
    _, r1 = db.calculate_quote(pid)
    db.add_price(a, 12000, "2026-09-01")
    _, r2 = db.calculate_quote(pid)
    assert r2.price > r1.price


def test_optional_switches(tmpdb):
    pid, _ = seed()
    s = db.get_settings()
    assert s["use_logistics"] is False and s["use_margin"] is False   # 기본: 선택 항목 꺼짐
    s.update(use_margin=True, margin_rate=.05)
    db.save_settings(s)
    q = db.build_quote_input(pid)
    assert q.margin_rate == .05 and q.logistics_rate == 0


def test_quote_snapshot_survives_price_change(tmpdb):
    pid, a = seed()
    q, r = db.calculate_quote(pid)
    qid = db.save_quote(q, r, "tester", "memo")
    db.add_price(a, 99999, "2026-12-01")
    q2, r2, row = db.load_quote(qid)
    assert r2.price == r.price and q2.ingredients[0].price_per_kg == 11000
    assert row["created_by"] == "tester"


def test_xlsx_export(tmpdb):
    pid, _ = seed()
    q, r = db.calculate_quote(pid)
    wb = load_workbook(BytesIO(quotes_to_xlsx([(q, r)], True)))
    assert wb.sheetnames[0] == "요약" and len(wb.sheetnames) == 2
    wb2 = load_workbook(BytesIO(quotes_to_xlsx([(q, r)], False)))
    assert wb2.sheetnames == ["납품가"]
    assert [c.value for c in wb2.active[2]] == ["테스트제품", r.price]   # 원가 정보 없음
